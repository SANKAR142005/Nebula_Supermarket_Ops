from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from app.operations.insights import sales_overview
from app.operations.purchasing import purchase_suggestions

def create_report(output='generated/store_report.pptx'):
    Path(output).parent.mkdir(parents=True,exist_ok=True); r=Presentation(); title=r.slides.add_slide(r.slide_layouts[0]); title.shapes.title.text='SmartStore Business Report'; title.placeholders[1].text='Inventory and sales snapshot'
    s=r.slides.add_slide(r.slide_layouts[1]); s.shapes.title.text='Sales'; box=s.placeholders[1].text_frame; box.text=str(sales_overview())
    s=r.slides.add_slide(r.slide_layouts[1]); s.shapes.title.text='Purchasing'; s.placeholders[1].text=str(purchase_suggestions())
    r.save(output); return {'success':True,'file':output}
